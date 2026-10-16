from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).with_name("EzRewards-UX-Audit-Report.docx")

PURPLE = "7C3AED"
PURPLE_DARK = "4C1D95"
PURPLE_LIGHT = "F3E8FF"
NAVY = "171229"
INK = "221D2F"
MUTED = "665F73"
LIGHT = "F7F5FA"
LINE = "DDD8E5"
WHITE = "FFFFFF"
RED = "B42318"
AMBER = "B54708"
GREEN = "067647"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_font(run, name, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def apply_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_cell_text(cell, bold=False, color=INK, size=9.2):
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        for run in p.runs:
            set_font(run, "Inter", size, bold=bold, color=color)


def add_table(doc, headers, rows, widths, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        hdr.cells[i].text = value
        shade_cell(hdr.cells[i], header_fill)
        style_cell_text(hdr.cells[i], bold=True, color=WHITE, size=9)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if r_idx % 2 == 1:
                shade_cell(cells[i], LIGHT)
            style_cell_text(cells[i])
    apply_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_font(r, "Inter", 10.5, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, "Inter", 10.5, color=INK)
    return p


def add_callout(doc, label, text, fill=PURPLE_LIGHT, accent=PURPLE_DARK):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    r = p.add_run(label.upper())
    set_font(r, "Inter", 8.5, bold=True, color=accent)
    r2 = p.add_run("\n" + text)
    set_font(r2, "Inter", 10.5, color=INK)


def add_screen_audit(doc, title, purpose, strengths, weaknesses, improvements, acceptance):
    doc.add_heading(title, level=2)
    p = doc.add_paragraph()
    r = p.add_run("User goal: ")
    set_font(r, "Inter", 10.5, bold=True, color=PURPLE_DARK)
    r = p.add_run(purpose)
    set_font(r, "Inter", 10.5, color=INK)
    doc.add_heading("What works", level=3)
    for item in strengths:
        add_bullet(doc, item)
    doc.add_heading("UX risks and weaknesses", level=3)
    for item in weaknesses:
        add_bullet(doc, item)
    doc.add_heading("Recommended improvements", level=3)
    for item in improvements:
        add_bullet(doc, item)
    add_callout(doc, "Developer acceptance criteria", acceptance, fill="F0FDF4", accent=GREEN)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.75)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(1.0)
sec.right_margin = Inches(1.0)
sec.header_distance = Inches(0.3)
sec.footer_distance = Inches(0.3)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Inter"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Inter")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Inter")
normal.font.size = Pt(10.5)
normal.font.color.rgb = rgb(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for name, size, color, before, after in (
    ("Heading 1", 17, PURPLE_DARK, 16, 8),
    ("Heading 2", 13.5, PURPLE_DARK, 12, 6),
    ("Heading 3", 11.2, NAVY, 8, 4),
):
    s = styles[name]
    s.font.name = "Space Grotesk"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Space Grotesk")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Space Grotesk")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = rgb(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

for section in doc.sections:
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = hp.add_run("EZREWARDS  |  UX AUDIT")
    set_font(rr, "Inter", 8, bold=True, color=MUTED)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("EzRewards UX Audit Report  •  Product & Engineering Edition")
    set_font(r, "Inter", 8, color=MUTED)

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(62)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("PRODUCT EXPERIENCE REVIEW")
set_font(r, "Inter", 9, bold=True, color=PURPLE)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run("EzRewards UX Audit Report")
set_font(r, "Space Grotesk", 30, bold=True, color=NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(24)
r = p.add_run("A flow-by-flow assessment of onboarding, dashboard, redemption, people, groups, and history experiences")
set_font(r, "Space Grotesk", 14, color=PURPLE_DARK)
add_callout(doc, "Executive takeaway", "The interface has a cohesive dark, purple-led identity and recognizable admin patterns. The largest usability risks are low-contrast supporting text, dense and modal-heavy forms, weak system-status feedback, and insufficient prevention and recovery around high-consequence actions.")
meta = [
    ("Audience", "Product management, design, frontend engineering, QA"),
    ("Audit basis", "Figma user-flow screenshots and the previously completed UX assessment"),
    ("Frameworks", "Nielsen's 10 usability heuristics, WCAG 2.2 AA principles, form and interaction design best practices"),
    ("Prepared", "20 August 2026"),
    ("Reference", "EzRewards website design - Figma node 91:244"),
]
add_table(doc, ["Report detail", "Value"], meta, [2200, 7160], header_fill=PURPLE_DARK)

doc.add_page_break()
doc.add_heading("Table of contents", level=1)
toc_items = [
    "1. Executive summary", "2. Scope, method, and limitations", "3. Cross-product findings",
    "4. Flow-by-flow screen audit", "5. Nielsen heuristic evaluation", "6. Accessibility and brand guidance",
    "7. Prioritized improvement roadmap", "8. Product and engineering acceptance criteria", "9. Measurement plan", "10. Appendix"
]
for item in toc_items:
    add_bullet(doc, item)
add_callout(doc, "Navigation note", "The document uses Word heading styles throughout, so readers can use the Navigation Pane and regenerate an automatic page-numbered table of contents in Word if the report changes.", fill=LIGHT, accent=PURPLE_DARK)

doc.add_heading("1. Executive summary", level=1)
doc.add_heading("Overall assessment", level=2)
doc.add_paragraph("EzRewards presents a promising, recognizable admin-console experience with a modern dark theme and a consistent purple accent. Core information architecture is familiar across dashboard, people, groups, redemption, and history. The design is visually cohesive, but the current screens frequently trade clarity for subtlety: labels, helper text, table separators, and inactive states are often too faint or too small, especially in dense workflows.")
doc.add_paragraph("The highest-risk moments are onboarding, redemption, and employee/group management. These flows require clear progress, robust validation, reversible actions, explicit confirmation, and strong feedback. The screenshots show several places where the system leaves users to infer what happened or what will happen next.")
add_callout(doc, "Recommended product decision", "Treat accessibility, form-state behavior, and system feedback as launch requirements, not polish. Simplify onboarding and invitation forms before expanding lower-priority dashboard embellishments.")

doc.add_heading("Strengths", level=2)
for x in [
    "Consistent dark navy and purple visual language creates a distinct product identity.",
    "Familiar admin navigation and content patterns reduce the learning curve.",
    "The main product areas map logically to common employer rewards tasks.",
    "Cards, tables, filters, and modals provide a workable component foundation.",
    "The flow structure anticipates common operational needs: invite people, organize groups, redeem rewards, and review history.",
]: add_bullet(doc, x)
doc.add_heading("Weaknesses", level=2)
for x in [
    "Low-contrast secondary text and fine dividers reduce readability and may fail WCAG AA.",
    "Typography is often small, particularly labels, helper copy, legal text, and table metadata.",
    "Loading, success, empty, validation, and recovery states are not consistently visible.",
    "Large forms and multi-purpose modals increase cognitive load and keyboard-navigation risk.",
    "High-consequence actions lack sufficient preview, confirmation, and persistent receipts.",
]: add_bullet(doc, x)

doc.add_heading("Priority summary", level=2)
priority_rows = [
    ("P0", "Accessibility and transaction safety", "Contrast, focus, keyboard, form validation, redemption review and receipt", "Blocks confident launch"),
    ("P1", "Flow clarity", "Simplify onboarding/invite; add loading, success, empty, and error states", "High conversion and support impact"),
    ("P2", "Information hierarchy", "Typography, tables, filters, responsive behavior", "Improves efficiency and comprehension"),
    ("P3", "System consistency", "Consolidate tokens, field patterns, status semantics, and motion", "Reduces future design and engineering debt"),
]
add_table(doc, ["Priority", "Theme", "Scope", "Why it matters"], priority_rows, [900, 1900, 3860, 2700])

doc.add_heading("2. Scope, method, and limitations", level=1)
doc.add_heading("Scope", level=2)
doc.add_paragraph("The audit covers the user-flow screenshots organized in the supplied Figma file: authentication and onboarding, dashboard, redemption, history, people management, and groups management. The report evaluates typography, color and accessibility, visual style, usability, and improvement opportunities for each flow.")
doc.add_heading("Method", level=2)
for x in [
    "Reviewed each screen in sequence to identify the user goal, decisions, system feedback, and potential failure points.",
    "Mapped findings to Nielsen's 10 usability heuristics and common UX audit principles: clarity, consistency, accessibility, efficiency, error prevention, recovery, and content hierarchy.",
    "Separated product-level recommendations from developer-ready behaviors and acceptance criteria.",
    "Assigned priority based on user harm, task failure risk, frequency, and implementation dependency.",
]: add_number(doc, x)
doc.add_heading("Limitations", level=2)
add_callout(doc, "Evidence limitation", "Several Figma frames were flattened screenshots. Exact font files, token values, computed contrast ratios, motion, responsive breakpoints, screen-reader output, keyboard order, and live validation behavior could not be verified from screenshots alone. Accessibility conformance must be confirmed in a coded build with automated and manual testing.", fill="FFF7ED", accent=AMBER)

doc.add_heading("3. Cross-product findings", level=1)
cross = [
    ("Typography", "Small and faint supporting text weakens hierarchy.", "Use Space Grotesk for headings and Inter for body/labels; body and inputs 16px minimum in product UI, labels 14px minimum, helper text 13-14px."),
    ("Color", "Purple-led branding is recognizable but overused for state and action.", "Reserve purple for primary/current state. Use distinct semantic colors plus text/icons. Verify WCAG contrast."),
    ("System feedback", "Many transitions and actions do not communicate progress or completion.", "Define loading, progress, success, empty, warning, error, and retry states for every async action."),
    ("Forms", "Dense two-column forms and modal forms raise completion cost.", "Group by intent, mark required fields, validate inline, preserve entered data, and use pages/side panels for long tasks."),
    ("Risky actions", "Consequences are not always previewed or recoverable.", "Add review steps, precise confirmation copy, reference IDs, undo where possible, and audit trails."),
    ("Responsive UX", "Wide tables and two-column layouts may not adapt cleanly.", "Define priority columns, card fallback, single-column forms, sticky actions, and 200% zoom behavior."),
]
add_table(doc, ["Area", "Finding", "Recommendation"], cross, [1300, 3000, 5060])

doc.add_heading("4. Flow-by-flow screen audit", level=1)
doc.add_heading("4.1 Onboarding flow", level=1)
add_screen_audit(doc, "Welcome / brand introduction", "Understand the product and begin setup.",
    ["Strong opportunity to establish the Visionary dark-theme identity.", "A focused first step can reduce distraction."],
    ["Sparse content may not explain what EzRewards does or how long setup will take.", "A decorative welcome state without a clear next action weakens task orientation."],
    ["State the value proposition in one sentence and list the setup outcome.", "Show estimated time and the number of steps.", "Provide one primary action and a clear sign-in path for returning users."],
    "A first-time user can identify the product purpose, expected setup time, current step, and next action without relying on imagery alone.")
add_screen_audit(doc, "Sign-in", "Enter credentials and continue into onboarding.",
    ["Email and password are familiar, low-learning-cost controls.", "A single primary action supports a focused task."],
    ["Labels, helper text, and legal copy appear small and low contrast.", "Password recovery, visibility, loading, and error behavior are not evident."],
    ["Use persistent labels above fields; do not rely on placeholders.", "Add show/hide password and password-recovery links.", "Validate on blur and submit, keep the user's email after errors, and move focus to an error summary when needed.", "For the prototype, allow any non-empty credentials and clearly label the experience as a demo if externally shared."],
    "Fields have programmatic labels, visible focus, accessible error text, password visibility control, loading state, and Enter-key submission. Any non-empty demo credentials route to the next step.")
add_screen_audit(doc, "Verification / transition state", "Understand that the account is being verified and know what to do next.",
    ["A dedicated transition screen can reduce uncertainty when verification is asynchronous."],
    ["The status, expected duration, destination email, and recovery options are unclear.", "Users may repeat actions if progress is not visible."],
    ["Use a descriptive title such as 'Check your inbox'.", "Show the masked destination email, expected delivery time, resend timer, and change-email action.", "Announce status changes to assistive technology."],
    "The screen explains what happened, where the message was sent, how long it may take, when resend becomes available, and how to correct the address.")
add_screen_audit(doc, "Workspace / profile setup", "Provide the minimum information needed to configure an organization.",
    ["The form captures operational context needed for a personalized workspace.", "Step-based setup can create a sense of progress."],
    ["Long uninterrupted forms increase fatigue and abandonment.", "Two-column form density can create ambiguous reading and tab order.", "Required and optional fields are not sufficiently differentiated."],
    ["Group fields into organization, administrator, and preferences sections.", "Ask only for information needed now; defer enrichment.", "Provide Back, Save and continue, and Save for later where appropriate.", "Use a single column on mobile and at 200% zoom; warn on abandonment only after edits."],
    "Every field has a visible label, required indicator, autocomplete where applicable, inline validation, logical DOM/tab order, persisted values when navigating back, and a clear progress indicator.")

doc.add_heading("4.2 Dashboard flow", level=1)
add_screen_audit(doc, "Dashboard overview", "Understand reward activity, balances, and items that need attention.",
    ["Sidebar and summary-card patterns are familiar.", "The dashboard can become a useful operational overview."],
    ["Content is horizontally spread and the primary insight is not obvious.", "Small supporting copy and weak chart hierarchy slow scanning.", "Freshness and date scope are not prominent."],
    ["Give each metric a clear label, primary value, comparison, and trend.", "Add chart axes, accessible summaries, tooltips, and a date-range control.", "Show data freshness and define loading, empty, error, and partial-data states.", "Indicate active navigation with more than color."],
    "A user can identify the reporting period, last refresh, most important metric, trend direction, and next action. Charts have text alternatives and keyboard-accessible data details.")

doc.add_heading("4.3 Redeem flow", level=1)
add_screen_audit(doc, "Redeem rewards", "Choose an amount, understand the exchange, and complete a redemption safely.",
    ["A dedicated redemption task supports a focused, high-value action.", "Balance-led context can reduce calculation effort."],
    ["Available, eligible, and entered amounts may be confused.", "Rules, fees, conversion, and processing time are not sufficiently visible before commitment.", "A single submit action risks accidental or duplicate redemption."],
    ["Show available balance, eligible balance, min/max, conversion rate, fees, and processing time before entry.", "Validate in real time and explain why the action is disabled.", "Add a review step with destination, amount, value, fees, and delivery estimate.", "After submission, show a persistent receipt with reference ID and history link."],
    "Invalid amounts cannot submit; disabled states include an explanation; the review screen repeats all material terms; submission is idempotent; success includes reference ID, timestamp, amount, destination, and status.")

doc.add_heading("4.4 History flow", level=1)
add_screen_audit(doc, "Redemption / transaction history", "Find, filter, understand, and reconcile past activity.",
    ["A tabular view is appropriate for repeatable transaction records.", "Search and filters support operational lookup."],
    ["Low-contrast separators and compact type reduce readability.", "Status meaning, date format, timezone, and no-result states are ambiguous.", "Wide tables may fail on small screens."],
    ["Group search, filters, date range, status, and amount controls.", "Show active filter chips, result count, and Reset filters.", "Pair status color with text and icon; define date format and timezone.", "Differentiate no history, no matches, loading, and error.", "Prioritize columns for narrow screens and provide an accessible detail view."],
    "Filters persist during navigation, results updates are announced, columns can be sorted with accessible state, dates include timezone context, and mobile users can reach every field without horizontal content loss.")

doc.add_heading("4.5 People flow", level=1)
add_screen_audit(doc, "People list and filtered state", "Find employees, understand status, and take permitted actions.",
    ["List, search, and filter patterns match user expectations.", "A centralized people view supports frequent administration."],
    ["Controls blend into the dark background.", "Filtered state and total results are not prominent.", "Row actions may be difficult to discover or operate by keyboard."],
    ["Strengthen boundaries and labels for search/filter controls.", "Show total count, applied filters, and a clear reset action.", "Define Active, Pending, Suspended, and Failed in plain language.", "Use an accessible action menu with named buttons and focus restoration."],
    "All list controls have names, keyboard access, visible focus, predictable order, and announced result counts. Status never relies on color alone.")
add_screen_audit(doc, "Edit person", "Update identity, employment, contact, and access details safely.",
    ["A dedicated edit experience can support deeper employee administration."],
    ["Mixed identity, HR, and permission fields make consequences difficult to understand.", "Unsaved state and permission-change effects are not explicit."],
    ["Separate Identity, Employment, Contact, and Access sections.", "Explain permission implications next to role controls.", "Show unsaved status, changed fields, and a precise save confirmation.", "Require re-authentication or elevated permission for sensitive changes if introduced later."],
    "Sensitive changes identify their consequences; only changed values are submitted; unsaved changes are protected; success states name what changed and for whom.")
add_screen_audit(doc, "Invite employee", "Invite a colleague with the minimum information required to start.",
    ["The proposed fields cover many future HR and organizational use cases."],
    ["The modal is too large and mixes invitation needs with optional HR enrichment.", "Two-column density and long keyboard paths increase error risk.", "Duplicate-email, manager, role, phone, and date validation are not clear."],
    ["Use a full page or side panel for the long version.", "Require only name, work email, and role to invite; defer optional details.", "Explain what email will be sent and when access begins.", "Validate duplicates and formats inline, and preserve values after recoverable errors.", "Warn on cancel only if the form is dirty."],
    "A valid invite can be completed with name, work email, and role. Duplicate emails are blocked with a recovery path; optional fields do not block submission; success confirms recipient, role, and resend/manage actions.")
add_screen_audit(doc, "Confirmation dialog", "Confirm or cancel a high-impact people action.",
    ["A confirmation step can prevent accidental destructive or permission-changing actions."],
    ["Generic confirmation copy does not communicate object, consequence, or reversibility.", "Focus management is not evident."],
    ["Name the person and action in the title and primary button.", "Explain consequences and whether undo is possible.", "Focus the safest logical control, trap focus, support Escape where safe, and restore focus to the trigger."],
    "The dialog has an accessible name and description, precise action label, logical initial focus, complete keyboard behavior, and focus restoration after close.")

doc.add_heading("4.6 Groups flow", level=1)
add_screen_audit(doc, "Groups list / empty state", "Understand existing groups or create the first one.",
    ["Groups provide a scalable way to organize reward programs and people."],
    ["A nearly empty screen can be mistaken for loading or failure.", "The purpose and downstream effect of a group are not explained."],
    ["Use an intentional empty state with a short explanation, example use cases, and Create group CTA.", "Show skeleton/loading separately from true emptiness.", "For populated lists, show purpose, owner, member count, and update time."],
    "Loading, empty, error, and populated states are visually and programmatically distinct. The empty state explains value and provides the permitted next action.")
add_screen_audit(doc, "Create group", "Name a group, explain its purpose, and select members.",
    ["Title, description, and membership are a reasonable minimum model."],
    ["Create and edit states may be visually similar.", "Comma-separated member entry is error-prone.", "Naming rules, duplicates, and permission implications are unclear."],
    ["Use a clear Create group title and button.", "Provide searchable multi-select with validated member chips and a live member count.", "State naming rules, detect duplicates, and explain what the group controls.", "Disable submission until required fields are valid and explain blockers."],
    "The UI distinguishes create from edit, prevents duplicate names where required, validates every member, exposes member count, and confirms the created group with a direct link.")
add_screen_audit(doc, "Group detail and edit", "Review group context and safely change its metadata or membership.",
    ["A detail view can centralize ownership, purpose, membership, and actions."],
    ["Row actions and membership changes may be hard to discover.", "The label 'Users' is less specific than 'Group members'.", "It is not clear who will be added or removed before save."],
    ["Show owner, purpose, member count, created date, and updated date.", "Use 'Group members' with searchable validated chips.", "Preview additions and removals and explain permission/reward implications.", "Enable Save changes only after an actual change and show persistent confirmation."],
    "Changed membership is summarized before save; additions/removals are explicit; errors do not discard edits; permissions are enforced server-side when a backend is introduced; success is persistent and auditable.")

doc.add_heading("5. Nielsen heuristic evaluation", level=1)
heuristics = [
    ("1", "Visibility of system status", "Weak", "Add progress, loading, freshness, filter-result counts, submission receipts, and persistent success/error states."),
    ("2", "Match between system and real world", "Fair", "Use plain-language reward, eligibility, role, employment, and group terminology; expose date/time context."),
    ("3", "User control and freedom", "Weak", "Provide Back, Cancel, Save for later, change email, reset filters, undo where safe, and dirty-form protection."),
    ("4", "Consistency and standards", "Fair", "Standardize field anatomy, modal behavior, button labels, status semantics, and confirmation patterns."),
    ("5", "Error prevention", "Weak", "Use constraints, inline validation, disabled-state rationale, review steps, duplicate protection, and idempotency."),
    ("6", "Recognition rather than recall", "Fair", "Keep labels persistent, surface rules in context, show current filters and balances, and avoid hidden consequences."),
    ("7", "Flexibility and efficiency", "Fair", "Add sorting, saved filters if needed, keyboard support, sensible defaults, and progressive disclosure."),
    ("8", "Aesthetic and minimalist design", "Good visually", "Preserve cohesion while increasing contrast; subtlety must not make content appear empty or disabled."),
    ("9", "Help users recover from errors", "Weak", "Use specific messages that name the problem, preserve input, recommend a fix, and provide retry/support paths."),
    ("10", "Help and documentation", "Weak", "Add contextual explanations for redemption rules, roles, groups, invitations, and account verification."),
]
add_table(doc, ["#", "Heuristic", "Assessment", "Required response"], heuristics, [500, 2450, 1250, 5160])

doc.add_heading("6. Accessibility and brand guidance", level=1)
doc.add_heading("Typography system", level=2)
type_rows = [
    ("Page title", "Space Grotesk", "24-28px / 600-700", "One primary page heading"),
    ("Section title", "Space Grotesk", "18-20px / 600", "Distinct section hierarchy"),
    ("Body and input", "Inter", "16px minimum / 400-500", "Readable at default zoom"),
    ("Label and table", "Inter", "14px minimum / 500-600", "Persistent, never placeholder-only"),
    ("Helper text", "Inter", "13-14px / 400", "Higher contrast than disabled text"),
    ("Button", "Inter", "14-16px / 600", "Action-led labels"),
]
add_table(doc, ["Role", "Font", "Recommended style", "Usage"], type_rows, [1600, 1700, 2200, 3860])
doc.add_heading("Color and contrast", level=2)
for x in [
    "Retain dark navy surfaces and purple accents as the core identity, but reserve saturated purple for primary actions and current selection.",
    "Meet WCAG AA contrast: at least 4.5:1 for normal text, 3:1 for large text, and 3:1 for meaningful UI boundaries and focus indicators.",
    "Increase contrast for placeholders, helper text, dividers, inactive navigation, disabled controls, table boundaries, and secondary icons.",
    "Pair success, warning, error, and status colors with text and/or icons; never communicate state with color alone.",
    "Use a 2px or stronger focus indicator with at least 3:1 contrast against adjacent colors.",
    "Use a modal backdrop that clearly separates layers while keeping the dialog boundary visible." ]: add_bullet(doc, x)
doc.add_heading("Interaction accessibility", level=2)
for x in [
    "Use semantic headings, landmarks, lists, tables, labels, buttons, dialogs, and live regions.",
    "Keep DOM order aligned with visual order; test keyboard-only completion for every flow.",
    "At 200% zoom, forms reflow to one column, actions remain available, and no content requires two-dimensional scrolling except true data tables.",
    "Touch targets should be at least 24 by 24 CSS pixels, with 44 by 44 preferred for frequent mobile actions.",
    "Respect reduced-motion preferences and avoid motion that is required to understand status." ]: add_bullet(doc, x)

doc.add_heading("7. Prioritized improvement roadmap", level=1)
roadmap = [
    ("P0", "Design system + frontend", "Contrast, type, focus, keyboard", "Token audit; focus styles; semantic field and dialog patterns", "All critical flows pass keyboard and WCAG AA checks"),
    ("P0", "Product + frontend", "Redemption safety", "Rules, real-time validation, review, idempotency, receipt", "No invalid/duplicate submit; receipt is durable"),
    ("P0", "Frontend + QA", "Form validation", "Accessible errors, summaries, preserved input, logical focus", "All required/error scenarios are test-covered"),
    ("P1", "Product + design", "Onboarding simplification", "Step model, progressive disclosure, back/save behavior", "Reduced fields and measurable completion funnel"),
    ("P1", "Product + design", "Invite simplification", "Name/email/role minimum; optional enrichment later", "Invite can complete quickly with clear outcome"),
    ("P1", "Frontend", "System states", "Loading, empty, success, error, retry, partial data", "Every async component has a defined state matrix"),
    ("P2", "Design + frontend", "Dashboard and table hierarchy", "Metrics, chart context, filters, responsive priorities", "Key information remains clear across breakpoints"),
    ("P2", "Frontend + QA", "Responsive and zoom", "Single-column forms; table fallback; sticky actions", "All flows pass 320px and 200% zoom tests"),
    ("P3", "Design system", "Pattern consolidation", "Tokens, statuses, buttons, fields, dialogs, notifications", "Reusable documented components replace one-offs"),
]
add_table(doc, ["Priority", "Owner", "Workstream", "Deliverable", "Definition of done"], roadmap, [700, 1400, 1700, 2800, 2760])

doc.add_heading("8. Product and engineering acceptance criteria", level=1)
doc.add_heading("Product manager checklist", level=2)
for x in [
    "Each flow has a defined user goal, entry point, success outcome, abandonment path, and analytics event.",
    "Required versus optional data is justified; later-stage enrichment does not block initial value.",
    "Rules and consequences are written in plain language before commitment.",
    "All status labels have agreed definitions and owner-approved copy.",
    "P0 and P1 findings are represented in the release plan with owners and measurable outcomes." ]: add_bullet(doc, x)
doc.add_heading("Developer checklist", level=2)
for x in [
    "Use native semantic elements first; custom controls expose accessible names, roles, values, and state.",
    "Focus order follows reading order; dialogs trap and restore focus; errors move focus predictably.",
    "Client validation improves feedback, while server validation remains authoritative when backend services are added.",
    "Submitting actions prevent duplicate requests and communicate loading without removing context.",
    "All form values survive recoverable errors and Back navigation within onboarding.",
    "Status, charts, filters, and results are understandable without color and operable without a mouse.",
    "Responsive tests cover narrow mobile, tablet, desktop, 200% browser zoom, and long localized strings." ]: add_bullet(doc, x)
doc.add_heading("QA evidence required", level=2)
qa_rows = [
    ("Accessibility", "Automated scan plus manual keyboard, screen-reader smoke test, contrast review, and 200% zoom"),
    ("Forms", "Required, invalid, duplicate, network error, retry, back navigation, dirty cancel, and successful submission"),
    ("Redemption", "Min/max, insufficient balance, conversion/fee display, double submit, success receipt, and failed/retried transaction"),
    ("Data views", "Loading, empty, no matches, error, partial data, sorting, filtering, pagination, and narrow viewport"),
    ("Dialogs", "Accessible name/description, initial focus, trap, Escape policy, action labels, and focus restoration"),
]
add_table(doc, ["Area", "Evidence"], qa_rows, [1900, 7460])

doc.add_heading("9. Measurement plan", level=1)
measure = [
    ("Onboarding", "Completion rate; step drop-off; median completion time; validation errors per session", "Confirms whether simplification improves activation"),
    ("Sign-in", "Success rate; recovery usage; error rate; repeated submit rate", "Surfaces credential and feedback problems"),
    ("Invite", "Invite completion; time to invite; duplicate-error rate; resend rate", "Measures administrative efficiency"),
    ("Redemption", "Successful redemption rate; abandoned review; invalid amount rate; duplicate attempts; support contacts", "Measures safety and confidence"),
    ("History", "Search/filter usage; no-result rate; time to locate transaction", "Measures findability"),
    ("Accessibility", "Keyboard task success; critical/serious issue count; contrast failures; zoom defects", "Tracks conformance and regression"),
]
add_table(doc, ["Flow", "Metrics", "Decision supported"], measure, [1500, 4800, 3060])

doc.add_heading("10. Appendix", level=1)
doc.add_heading("Severity model", level=2)
severity = [
    ("P0 - Critical", "Can block task completion, cause financial/permission mistakes, or prevent access for users with disabilities."),
    ("P1 - High", "Creates significant confusion, abandonment, support burden, or repeated errors."),
    ("P2 - Medium", "Reduces efficiency, comprehension, or responsive usability but has a workaround."),
    ("P3 - Low / systemic", "Consistency and maintainability improvement with lower immediate user harm."),
]
add_table(doc, ["Priority", "Definition"], severity, [1900, 7460])
doc.add_heading("Audit principles applied", level=2)
for x in [
    "Clarity and hierarchy", "Consistency and standards", "Accessibility and inclusion", "Task efficiency",
    "Error prevention and recovery", "System feedback and trust", "Progressive disclosure", "Responsive resilience", "Content usefulness", "Measurability"
]: add_bullet(doc, x)
doc.add_heading("Source reference", level=2)
doc.add_paragraph("Figma: EzRewards website design, node 91:244. URL supplied by the project owner. This report records the earlier screenshot-based audit and converts it into a product and engineering implementation brief.")
add_callout(doc, "Final note", "This document is an expert review, not a substitute for moderated usability testing. Validate the highest-risk assumptions with representative administrators and reward recipients before broad release.", fill=LIGHT, accent=PURPLE_DARK)

doc.core_properties.title = "EzRewards UX Audit Report"
doc.core_properties.subject = "Product and engineering UX audit"
doc.core_properties.author = "EzRewards Product Team"
doc.core_properties.keywords = "EzRewards, UX audit, onboarding, accessibility, heuristics"
doc.save(OUT)
print(OUT)
